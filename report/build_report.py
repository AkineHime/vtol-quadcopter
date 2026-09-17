# -*- coding: utf-8 -*-
"""Build the BCSE497J Project-I Review-2 report from template.docx.

    uv run --with python-docx python build_report.py

Reads content.py for all prose. Deletes the template's cyber-security sample
blocks and guidance notes, fills the title page, every section 1-5, fixes the
template's numbering typos, and inserts the seven figures. Output: report.docx
"""
from pathlib import Path

from docx import Document
from docx.document import Document as _Doc
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

import content as C

HERE = Path(__file__).resolve().parent
FIGS = HERE / "figs"
R2FIGS = HERE.parent / "documents" / "review2_figures"
SRC = HERE / "template.docx"
OUT = HERE / "report.docx"

FONT = "Times New Roman"


# --------------------------------------------------------------- block walking
def iter_blocks(parent):
    el = parent.element.body if isinstance(parent, _Doc) else parent._tc
    for child in el.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def blocks(doc):
    return list(iter_blocks(doc))


def find_para(doc, text, contains=False):
    for b in iter_blocks(doc):
        if isinstance(b, Paragraph):
            t = b.text.strip()
            if (text in t) if contains else (t == text or t.rstrip() == text):
                return b
    raise KeyError(text)


# ------------------------------------------------------------- para/run helpers
def _style_run(r, size=12, bold=False, italic=False, color=None):
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = color


def mk_para(doc, text="", size=12, bold=False, italic=False, align=AL.JUSTIFY,
            line=1.15, before=0, after=6, color=None):
    p = doc.add_paragraph()          # created at end of body; moved into place
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if text:
        _style_run(p.add_run(text), size, bold, italic, color)
    return p


def place_after(anchor, block):
    """Move a freshly made paragraph/table element to sit after `anchor`."""
    a = anchor._p if isinstance(anchor, Paragraph) else anchor._tbl
    b = block._p if isinstance(block, Paragraph) else block._tbl
    a.addnext(b)
    return block


def run_seq(doc, anchor, builders):
    """builders: list of callables(doc)->block. Insert them in order after
    anchor, advancing the cursor."""
    cur = anchor
    for build in builders:
        blk = build(doc)
        place_after(cur, blk)
        cur = blk
    return cur


def delete_block(b):
    el = b._p if isinstance(b, Paragraph) else b._tbl
    el.getparent().remove(el)


# ---------------------------------------------------------------- content bits
def para(text, **kw):
    return lambda doc: mk_para(doc, text, **kw)


def heading(text, level):
    sz = {1: 14, 2: 13, 3: 12}[level]
    al = AL.CENTER if level == 1 else AL.LEFT
    return lambda doc: mk_para(doc, text, size=sz, bold=True,
                               italic=(level == 3), align=al, line=1.08,
                               before=6, after=8)


def bullets(items, sub=False):
    def build_all(doc):
        return items  # not used; handled specially
    return items


def fig(path, caption, width=6.1):
    def build(doc):
        p = mk_para(doc, "", align=AL.CENTER, after=2, before=8)
        p.add_run().add_picture(str(path), width=Inches(width))
        return p
    cap = lambda doc: mk_para(doc, caption, align=AL.CENTER, size=11,
                              italic=True, after=10, before=0)
    return [build, cap]


def kv_table(doc, rows, w0=2.0, w1=4.2, head=None):
    t = doc.add_table(rows=len(rows) + (1 if head else 0), cols=2)
    t.style = "Table Grid"
    t.autofit = False
    ri = 0
    if head:
        for ci, htext in enumerate(head):
            c = t.rows[0].cells[ci]
            c.text = ""
            _style_run(c.paragraphs[0].add_run(htext), 11, bold=True)
        ri = 1
    for (a, b) in rows:
        for ci, val in enumerate((a, b)):
            c = t.rows[ri].cells[ci]
            c.text = ""
            _style_run(c.paragraphs[0].add_run(val), 11,
                       bold=(ci == 0 and w0 < 1.6))
        ri += 1
    for row in t.rows:
        row.cells[0].width = Inches(w0)
        row.cells[1].width = Inches(w1)
    return t


def num_list(doc, anchor, items, tag_bold=True):
    """items: list of (tag, body) or (tag, title, body)."""
    cur = anchor
    for it in items:
        if len(it) == 3:
            tag, title, body = it
            txt = None
        else:
            tag, body = it
            title = None
        p = mk_para(doc, "", after=6)
        _style_run(p.add_run(f"{tag}. "), 12, bold=True)
        if title:
            _style_run(p.add_run(f"{title} — "), 12, bold=True)
        _style_run(p.add_run(body), 12)
        place_after(cur, p)
        cur = p
    return cur


# =============================================================== BUILD
def main():
    doc = Document(str(SRC))

    # ---- 1. delete sample ranges (descending so indices stay valid) --------
    bl = blocks(doc)
    ranges = [(31, 58), (64, 75), (91, 98), (103, 109), (114, 120),
              (155, 156), (167, 175), (179, 186), (190, 207), (211, 230),
              (238, 254), (267, 281)]
    for lo, hi in sorted(ranges, reverse=True):
        for i in range(hi, lo - 1, -1):
            delete_block(bl[i])

    # ---- 2. delete guidance / formatting-note paragraphs -------------------
    KILL_PREFIX = (
        "(Times New Roman", "<Contents, Times New Roman", "<Contents , Times",
        "(mention the particulars", "(Mention the particulars",
        "(Line spacing", "(Specialization if any)", "Remove Border in Table)",
        "One page and not exceeding", "< This section of a project report",
        "< Maximum of 50 journal", "<To provide the reader",
        "< To explain why the project", "<To define the boundaries",
        "< To clearly define what the project", "< Clearly define the issue",
        "<A Gantt chart helps", "<To identify the areas",
        "(mention the particular", "< To provide the reader",
    )
    for b in list(iter_blocks(doc)):
        if isinstance(b, Paragraph):
            t = b.text.strip()
            if t.startswith(KILL_PREFIX):
                delete_block(b)

    # ---- 3. title page ----------------------------------------------------
    tp = find_para(doc, "<Title of Project>")
    tp.text = ""
    _style_run(tp.add_run(C.TITLE), 16, bold=True)
    tp.alignment = AL.CENTER

    tbls = [b for b in iter_blocks(doc) if isinstance(b, Table)]
    team_tbl = next(t for t in tbls
                    if "Reg. No." in t.rows[0].cells[0].text)
    for i, (reg, name) in enumerate(C.TEAM):
        for ci, (val, bold) in enumerate(((reg, False), (name.upper(), True))):
            cell = team_tbl.rows[i].cells[ci]
            cell.text = ""
            _style_run(cell.paragraphs[0].add_run(val), 12, bold=bold)
    guide_tbl = next(t for t in tbls
                     if "Project Guide Name" in t.rows[0].cells[0].text)
    gc = guide_tbl.rows[0].cells[0]
    align0 = gc.paragraphs[0].alignment
    gc.text = ""
    gc.paragraphs[0].alignment = align0 or AL.CENTER
    _style_run(gc.paragraphs[0].add_run(C.GUIDE_NAME), 12, bold=True)

    # ---- 4. ABSTRACT ----------------------------------------------------
    run_seq(doc, find_para(doc, "ABSTRACT"),
            [para(t) for t in C.ABSTRACT] + [
                lambda doc: mk_para(
                    doc,
                    "Keywords — autonomous UAV; VTOL quadplane; ArduPilot; "
                    "coconut disease and pest detection; YOLO object "
                    "detection; flight simulation.", italic=True, after=6)])

    # ---- 5. INTRODUCTION ------------------------------------------------
    run_seq(doc, find_para(doc, "1.1 Background", contains=True),
            [para(t) for t in C.BACKGROUND])
    run_seq(doc, find_para(doc, "1.2 Motivation", contains=True),
            [para(t) for t in C.MOTIVATION])
    run_seq(doc, find_para(doc, "1.3 Scope of the Project", contains=True),
            [para(t) for t in C.SCOPE])

    # ---- 6. LITERATURE REVIEW + table ---------------------------------
    anchor = find_para(doc, "2.1 Literature Review", contains=True)
    cur = anchor
    for t in C.LIT_INTRO:
        cur = place_after(cur, mk_para(doc, t))
    heads = ("Paper", "Objective", "Methodology", "Pros / Cons", "Findings")
    widths = (1.15, 1.55, 1.6, 1.35, 1.15)
    lt = doc.add_table(rows=len(C.LIT_TABLE) + 1, cols=5)
    lt.style = "Table Grid"
    lt.autofit = False
    for ci, h in enumerate(heads):
        c = lt.rows[0].cells[ci]
        c.text = ""
        _style_run(c.paragraphs[0].add_run(h), 9, bold=True)
    for ri, row_vals in enumerate(C.LIT_TABLE, start=1):
        for ci, val in enumerate(row_vals):
            c = lt.rows[ri].cells[ci]
            c.text = ""
            _style_run(c.paragraphs[0].add_run(val), 8.5)
    for row in lt.rows:
        for ci, w in enumerate(widths):
            row.cells[ci].width = Inches(w)
    place_after(cur, lt)
    cur = place_after(lt, mk_para(
        doc, "Table 1.  Reviewed literature.", align=AL.CENTER, size=11,
        italic=True, after=10))
    for t in C.LIT_AFTER:
        cur = place_after(cur, mk_para(doc, t))

    # ---- 7. RESEARCH GAP ---------------------------------------------
    run_seq(doc, find_para(doc, "2.2 Research Gap", contains=True),
            [para(t) for t in C.RESEARCH_GAP])

    # ---- 8. OBJECTIVES ---------------------------------------------
    anchor = find_para(doc, "2.3 Objectives", contains=True)
    cur = place_after(anchor, mk_para(
        doc, "The project has three objectives, each stated so that it can be "
             "checked against a concrete deliverable."))
    cur = num_list(doc, cur, C.OBJECTIVES)
    for t in C.OBJ_TAIL:
        cur = place_after(cur, mk_para(doc, t))

    # ---- 9. PROBLEM STATEMENT ------------------------------------
    run_seq(doc, find_para(doc, "2.4 Problem Statement", contains=True),
            [para(t) for t in C.PROBLEM_STATEMENT])

    # ---- 10. PROJECT PLAN + Gantt ------------------------------------
    anchor = find_para(doc, "2.5 Project Plan", contains=True)
    cur = anchor
    for t in C.PROJECT_PLAN:
        cur = place_after(cur, mk_para(doc, t))
    # Gantt image goes right before the existing "Fig. 1. Gantt chart" caption
    gcap = find_para(doc, "Fig. 1. Gantt chart", contains=True)
    gimg = mk_para(doc, "", align=AL.CENTER, after=2, before=8)
    gimg.add_run().add_picture(str(FIGS / "gantt.png"), width=Inches(6.5))
    gcap._p.addprevious(gimg._p)
    gcap.alignment = AL.CENTER
    for r in gcap.runs:
        _style_run(r, 11, italic=True)

    # ---- 11. 3.1.1 FUNCTIONAL ------------------------------------
    anchor = find_para(doc, "Functional", contains=True)  # 3.1.1
    cur = place_after(anchor, mk_para(
        doc, "The functional requirements below define what the complete "
             "system must do; those marked for Project-II are specified now "
             "but demonstrated later."))
    cur = num_list(doc, cur, C.FUNCTIONAL)

    # ---- 12. 3.1.2 NON-FUNCTIONAL -------------------------------
    anchor = find_para(doc, "Non-Functional", contains=True)
    cur = anchor
    for name, body in C.NONFUNCTIONAL:
        p = mk_para(doc, "", after=5)
        _style_run(p.add_run(f"{name}. "), 12, bold=True)
        _style_run(p.add_run(body), 12)
        cur = place_after(cur, p)

    # ---- 13. 3.2 FEASIBILITY (create 3.2.1/2/3 sub-headings) --------
    anchor = find_para(doc, "3.2 Feasibility Study", contains=True)
    builders = [heading("3.2.1 Technical Feasibility", 3)]
    builders += [para(t) for t in C.TECH_FEAS]
    cur = run_seq(doc, anchor, builders)
    # experiment figures + key numbers inside technical feasibility
    for pth, cap, w in [
        (R2FIGS / "1_sitl_mission.png",
         "Fig. 5.  ArduPilot SITL — full autonomous mission; every "
         "flight-phase check passes.", 6.4),
        (R2FIGS / "2_aero_coefficients.png",
         "Fig. 6.  Real 3-D aerodynamic coefficients of the SD7037 wing / "
         "NACA 0009 tail (NeuralFoil + AeroSandbox).", 6.4),
        (R2FIGS / "3_bridge_survey_instability.png",
         "Fig. 7.  Full mission on the higher-fidelity aerodynamics — "
         "the survey-phase instability under diagnosis.", 5.8)]:
        for build in fig(pth, cap, w):
            cur = place_after(cur, build(doc))
    kt = kv_table(doc, C.KEY_NUMBERS, w0=3.1, w1=3.1,
                  head=("Metric", "Result (simulation / analysis)"))
    place_after(cur, kt)
    cur = place_after(kt, mk_para(
        doc, "Table 2.  Key simulation and analysis results to date.",
        align=AL.CENTER, size=11, italic=True, after=10))
    cur = run_seq(doc, cur, [heading("3.2.2 Economic Feasibility", 3)]
                  + [para(t) for t in C.ECON_FEAS]
                  + [heading("3.2.3 Social Feasibility", 3)]
                  + [para(t) for t in C.SOCIAL_FEAS])

    # ---- 14. 3.3 SYSTEM SPECIFICATION (fix 3.2 -> 3.3) --------------
    ssp = find_para(doc, "3.2 System Specification")
    ssp.text = ""
    _style_run(ssp.add_run("3.3 System Specification"), 13, bold=True)
    ssp.alignment = AL.LEFT
    cur = run_seq(doc, ssp, [
        para("The hardware and software actually required by the project are "
             "listed below."),
        heading("3.3.1 Hardware Specification", 3)])
    ht = kv_table(doc, C.HARDWARE_ROWS, w0=1.7, w1=4.6,
                  head=("Item", "Specification"))
    place_after(cur, ht)
    cur = place_after(ht, mk_para(
        doc, "Table 3.  Hardware specification (planned build; not purchased "
             "in Project-I).", align=AL.CENTER, size=11, italic=True,
        after=10))
    cur = place_after(cur, heading("3.3.2 Software Specification", 3)(doc))
    st = kv_table(doc, C.SOFTWARE_ROWS, w0=1.9, w1=4.4,
                  head=("Layer", "Tool / approach"))
    place_after(cur, st)
    place_after(st, mk_para(
        doc, "Table 4.  Software stack.", align=AL.CENTER, size=11,
        italic=True, after=10))

    # ---- 15. 4.1 SYSTEM ARCHITECTURE --------------------------------
    anchor = find_para(doc, "4.1 System Architecture", contains=True)
    cur = anchor
    aimg = mk_para(doc, "", align=AL.CENTER, after=2, before=8)
    aimg.add_run().add_picture(str(FIGS / "architecture.png"),
                               width=Inches(6.5))
    cur = place_after(cur, aimg)
    cur = place_after(cur, mk_para(
        doc, "Fig. 2.  System architecture — aerial tier and ground "
             "(laptop) tier.", align=AL.CENTER, size=11, italic=True,
        after=10))
    for t in C.ARCH_TEXT:
        cur = place_after(cur, mk_para(doc, t))

    # ---- 16. 4.2 DESIGN : DFD + Use Case + optional notes ----------
    dfd_h = find_para(doc, "4.2.1 Data Flow Diagram", contains=True)
    for _h, _t in [(dfd_h, "4.2.1 Data Flow Diagram")]:
        _h.text = ""
        _style_run(_h.add_run(_t), 12, bold=True, italic=True)
    cur = dfd_h
    dimg = mk_para(doc, "", align=AL.CENTER, after=2, before=8)
    dimg.add_run().add_picture(str(FIGS / "dfd.png"), width=Inches(6.5))
    cur = place_after(cur, dimg)
    cur = place_after(cur, mk_para(
        doc, "Fig. 3.  Level-1 data flow diagram.", align=AL.CENTER,
        size=11, italic=True, after=10))
    for t in C.DFD_TEXT:
        cur = place_after(cur, mk_para(doc, t))

    uc_h = find_para(doc, "4.2.2 Use Case Diagram", contains=True)
    uc_h.text = ""
    _style_run(uc_h.add_run("4.2.2 Use Case Diagram"), 12, bold=True,
               italic=True)
    cur = uc_h
    uimg = mk_para(doc, "", align=AL.CENTER, after=2, before=8)
    uimg.add_run().add_picture(str(FIGS / "usecase.png"), width=Inches(6.5))
    cur = place_after(cur, uimg)
    cur = place_after(cur, mk_para(
        doc, "Fig. 4.  Use-case diagram.", align=AL.CENTER, size=11,
        italic=True, after=10))
    for t in C.USECASE_TEXT:
        cur = place_after(cur, mk_para(doc, t))

    # replace 4.2.3 / 4.2.4 one-liners with the "not included" explanation
    for needle, txt in (("4.2.3 Class Diagram", C.CLASS_SEQ_TEXT[0]),
                        ("4.2.4 Sequence Diagram", C.CLASS_SEQ_TEXT[1])):
        p = find_para(doc, needle, contains=True)
        p.text = ""
        p.alignment = AL.JUSTIFY
        lead, _, rest = txt.partition(" — ")
        _style_run(p.add_run(lead + " — "), 12, bold=True, italic=True)
        _style_run(p.add_run(rest), 12)
    try:
        delete_block(find_para(
            doc, "Please include whichever diagram is applicable",
            contains=True))
    except KeyError:
        pass

    # ---- 17. REFERENCES ------------------------------------------
    anchor = find_para(doc, "5. REFERENCES", contains=True)
    cur = anchor
    n = 0
    for cat, items in C.REFS.items():
        h = mk_para(doc, "", after=4, before=8)
        parts = cat.split("<")
        _style_run(h.add_run(parts[0].strip()), 12, bold=True)
        if len(parts) > 1:
            _style_run(h.add_run("  <" + parts[1]), 11, italic=True)
        cur = place_after(cur, h)
        for ref in items:
            if ref.startswith("["):        # editorial note, keep verbatim
                cur = place_after(cur, mk_para(doc, ref, after=4))
            else:
                n += 1
                cur = place_after(cur, mk_para(doc, f"[{n}] {ref}", after=4))

    # ---- 18. TOC: strip helper hints, fix typos --------------------
    for b in iter_blocks(doc):
        if isinstance(b, Table) and b.rows and b.rows[0].cells[0].text.strip() \
                == "Sl.No":
            for row in b.rows:
                c1 = row.cells[1]
                full = c1.text
                new = full
                if "<" in new:
                    new = new[:new.index("<")].rstrip()
                new = new.replace("3.2.2 Social", "3.2.3 Social")
                new = new.replace("1.2 Motivations", "1.2 Motivation")
                if new != full:
                    runs = c1.paragraphs[0].runs
                    if runs:
                        runs[0].text = new
                        for rr in runs[1:]:
                            rr.text = ""

    # ---- 19. collapse blank-paragraph runs in the body ------------
    def _is_blank(p):
        return isinstance(p, Paragraph) and not p.text.strip() and \
            p._p.find(qn("w:pPr")) is None or (
                isinstance(p, Paragraph) and not p.text.strip()
                and p._p.find(qn(".//w:drawing")) is None)

    def _has_pagebreak(p):
        return 'w:br' in p._p.xml and 'type="page"' in p._p.xml

    seq = list(iter_blocks(doc))
    # locate TOC
    toc_i = next(i for i, b in enumerate(seq) if isinstance(b, Paragraph)
                 and b.text.strip() == "TABLE OF CONTENTS")
    # (a) exactly one blank/pagebreak paragraph right before the TOC
    j = toc_i - 1
    blanks_before = []
    while j >= 0 and isinstance(seq[j], Paragraph) and not seq[j].text.strip():
        blanks_before.append(seq[j])
        j -= 1
    if len(blanks_before) > 1:
        for p in blanks_before[1:]:
            delete_block(p)
    # (b) after the TOC: collapse runs of >=2 blank paras (no image, no
    #     pagebreak) down to one
    after = list(iter_blocks(doc))
    start = next(i for i, b in enumerate(after) if isinstance(b, Paragraph)
                 and b.text.strip() == "TABLE OF CONTENTS")
    run = []
    for b in after[start + 1:]:
        blank = (isinstance(b, Paragraph) and not b.text.strip()
                 and not _has_pagebreak(b)
                 and b._p.find(qn("w:drawing")) is None
                 and not b._p.findall(".//" + qn("w:drawing")))
        if blank:
            run.append(b)
        else:
            for p in run[1:]:
                delete_block(p)
            run = []
    for p in run[1:]:
        delete_block(p)

    doc.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
