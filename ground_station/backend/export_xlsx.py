"""The ripeness prediction as an Excel workbook (GET /api/prediction.xlsx).

Sheets: Summary (counts as formulas over the Prediction sheet, the threshold as
an editable input), Prediction (one row per palm, most likely tender first),
Next mission (the planner's stops in flying order), How it works (the rules)."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import ripeness

MATURITY = {"unknown": "Not seen yet", "young": "Young", "tender": "Tender", "mature": "Mature",
            "over-mature": "Over-mature"}
HEALTH = {"healthy": "Healthy", "frond-damage": "Frond damage", "pest": "Pest"}
STATUS = {"unscouted": "Unscouted", "target": "Target", "confirmed": "Confirmed", "harvested": "Harvested"}

FONT = "Arial"
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F4F3A")
TITLE = Font(name=FONT, bold=True, size=14)
BOLD = Font(name=FONT, bold=True)
BODY = Font(name=FONT)
NOTE = Font(name=FONT, italic=True, color="666666", size=9)
INPUT = Font(name=FONT, color="0000FF", bold=True)          # blue = editable input
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
STAGE_FILL = {"Tender": "C6EFCE", "Young": "E2F0D9", "Mature": "FFF2CC", "Over-mature": "E8D5C4"}


def _when(iso: str | None):
    """ISO-8601 UTC -> naive datetime (Excel has no time zones; headers say UTC)."""
    if not iso:
        return None
    return datetime.fromisoformat(iso).astimezone(timezone.utc).replace(tzinfo=None)


def _header(ws, row: int, names: list[str], widths: list[float]) -> None:
    for i, (n, w) in enumerate(zip(names, widths), start=1):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def workbook(trees: list[dict], plan: dict, farm: dict | None = None) -> bytes:
    """trees: list_trees rows plus ripe_score / ripe_reason; plan: planner.public()."""
    now = datetime.now(timezone.utc)
    stops = {s["tree"]: s for s in plan.get("stops", [])}
    dropped = set(plan.get("dropped", []))
    rows = sorted(trees, key=lambda t: (-(t.get("ripe_score") or 0.0), t["id"]))
    n = len(rows)
    last = n + 1                                   # last data row on the Prediction sheet

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    pred = wb.create_sheet("Prediction")
    mission = wb.create_sheet("Next mission")
    about = wb.create_sheet("How it works")

    # ---------- Prediction: one row per palm ----------
    cols = ["Palm", "Row", "Column", "Tender probability", "Predicted tender", "Why", "Maturity", "Health",
            "Status", "Next mission", "Stop #", "Stop type", "Last visit (UTC)", "Last detection label",
            "Detection confidence", "Latitude", "Longitude"]
    _header(pred, 1, cols, [9, 6, 8, 12, 11, 38, 13, 13, 12, 14, 8, 11, 18, 18, 12, 12, 12])
    for r, t in enumerate(rows, start=2):
        s = stops.get(t["id"])
        in_mission = "Yes" if s else ("Left out (battery)" if t["id"] in dropped else "No")
        vals = [t["id"], t["row"], t["col"], t.get("ripe_score"),
                f'=IF(D{r}>=Summary!$C$5,"Yes","No")',
                t.get("ripe_reason"), MATURITY.get(t["maturity"], t["maturity"]), HEALTH.get(t["health"], t["health"]),
                STATUS.get(t["status"], t["status"]), in_mission, s["order"] if s else None,
                s["kind"].title() if s else None, _when(t.get("last_visit")), t.get("label"), t.get("confidence"),
                t["lat"], t["lon"]]
        for c, v in enumerate(vals, start=1):
            cell = pred.cell(row=r, column=c, value=v)
            cell.font = BODY
        pred.cell(row=r, column=4).number_format = "0%"
        pred.cell(row=r, column=13).number_format = "yyyy-mm-dd hh:mm"
        pred.cell(row=r, column=15).number_format = "0%"
        pred.cell(row=r, column=16).number_format = "0.000000"
        pred.cell(row=r, column=17).number_format = "0.000000"
        fill = STAGE_FILL.get(vals[6])
        if fill:
            pred.cell(row=r, column=7).fill = PatternFill("solid", fgColor=fill)
    if n:
        pred.conditional_formatting.add(f"D2:D{last}", ColorScaleRule(
            start_type="num", start_value=0, start_color="F8F8F8",
            mid_type="num", mid_value=0.5, mid_color="FFE699",
            end_type="num", end_value=1, end_color="63BE7B"))
    pred.freeze_panes = "B2"
    pred.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    # ---------- Summary: inputs + formulas over the Prediction sheet ----------
    summary.column_dimensions["A"].width = 2
    summary.column_dimensions["B"].width = 34
    summary.column_dimensions["C"].width = 14
    summary.column_dimensions["D"].width = 52
    summary["B1"] = "Tender-nut prediction"
    summary["B1"].font = TITLE
    farm_txt = (f"{n} palms" + (f" · {farm['grid']} x {farm['grid']} grid, {farm['spacing']} m spacing" if farm else ""))
    summary["B2"] = f"{farm_txt} · exported {now:%Y-%m-%d %H:%M} UTC from the ground station"
    summary["B2"].font = NOTE

    summary["B4"] = "Setting"
    summary["B5"] = "Predicted-tender threshold"
    summary["C5"] = ripeness.PREDICT_MIN
    summary["C5"].font, summary["C5"].fill, summary["C5"].number_format = INPUT, INPUT_FILL, "0%"
    summary["D5"] = "Edit to try another cut-off: the Yes/No column and the counts below update."
    summary["C5"].comment = Comment("Same value as PREDICT_MIN in backend/ripeness.py (the dashboard and planner "
                                    "use it). Changing it here only changes this workbook.", "ground station")

    P = "Prediction!"
    rng = lambda col: f"{P}${col}$2:${col}${last}"          # noqa: E731
    lines = [
        ("Overview", None, None),
        ("Palms", f"=COUNTA({rng('A')})", None),
        ("Predicted tender (at or above the threshold)", f'=COUNTIF({rng("E")},"Yes")', None),
        ("Share of the farm predicted tender", "=IF(C8=0,0,C9/C8)", "0%"),
        ("Average tender probability", f"=IF(C8=0,0,AVERAGE({rng('D')}))", "0%"),
        ("Palms in the next mission", f'=COUNTIF({rng("J")},"Yes")', None),
        ("Left out by the battery limit", f'=COUNTIF({rng("J")},"Left out (battery)")', None),
        ("Tender probability", None, None),
        ("80% and above", f'=COUNTIFS({rng("D")},">=0.8")', None),
        ("50% to 80%", f'=COUNTIFS({rng("D")},">=0.5",{rng("D")},"<0.8")', None),
        ("35% to 50%", f'=COUNTIFS({rng("D")},">=0.35",{rng("D")},"<0.5")', None),
        ("Below 35%", f'=COUNTIFS({rng("D")},"<0.35")', None),
        ("Maturity (last sighting)", None, None),
        *[(lab, f'=COUNTIF({rng("G")},"{lab}")', None) for lab in MATURITY.values()],
        ("Health", None, None),
        *[(lab, f'=COUNTIF({rng("H")},"{lab}")', None) for lab in HEALTH.values()],
        ("Status", None, None),
        *[(lab, f'=COUNTIF({rng("I")},"{lab}")', None) for lab in STATUS.values()],
    ]
    r = 7
    for label, formula, fmt in lines:
        if formula is None:
            r += 1 if r > 7 else 0
            summary.cell(row=r, column=2, value=label).font = BOLD
            summary.cell(row=r, column=2).border = THIN
        else:
            summary.cell(row=r, column=2, value=label).font = BODY
            c = summary.cell(row=r, column=3, value=formula)
            c.font = BODY
            if fmt:
                c.number_format = fmt
        r += 1

    r += 1
    summary.cell(row=r, column=2, value="Next mission (from the planner)").font = BOLD
    summary.cell(row=r, column=2).border = THIN
    r += 1
    mission_lines = [
        ("Hover stops", len(plan.get("stops", [])), None),
        ("  targets (marked by a person)", plan.get("n_targets"), None),
        ("  predicted tender", plan.get("n_predicted"), None),
        ("Route length (km)", round((plan.get("distance_m") or 0) / 1000, 2), "0.00"),
        ("Energy (Wh)", plan.get("energy_wh"), "0.0"),
        ("Battery used", (plan.get("battery_pct") or 0) / 100, "0%"),
        ("Battery limit", (plan.get("budget_pct") or 0) / 100, "0%"),
    ]
    for label, v, fmt in mission_lines:
        summary.cell(row=r, column=2, value=label).font = BODY
        c = summary.cell(row=r, column=3, value=v)
        c.font = BODY
        if fmt:
            c.number_format = fmt
        r += 1
    summary.cell(row=r, column=2, value=f"Planner values at export time ({plan.get('generated_at', '')[:16].replace('T', ' ')} UTC): "
                 "44.4 Wh pack, 7.3 Wh sweep + 0.9 Wh return + 1.2 Wh per stop.").font = NOTE

    # ---------- Next mission: stops in flying order ----------
    _header(mission, 1, ["Order", "Palm", "Type", "Tender probability", "Why", "Latitude", "Longitude"],
            [7, 9, 11, 12, 40, 12, 12])
    for r, s in enumerate(sorted(plan.get("stops", []), key=lambda s: s["order"]), start=2):
        for c, v in enumerate([s["order"], s["tree"], s["kind"].title(), s.get("score"), s.get("reason"),
                               s["lat"], s["lon"]], start=1):
            mission.cell(row=r, column=c, value=v).font = BODY
        mission.cell(row=r, column=4).number_format = "0%"
        mission.cell(row=r, column=6).number_format = "0.000000"
        mission.cell(row=r, column=7).number_format = "0.000000"
    mission.freeze_panes = "A2"
    if dropped:
        r = len(plan.get("stops", [])) + 3
        mission.cell(row=r, column=1, value="Left out (battery): " + ", ".join(sorted(dropped))).font = NOTE

    # ---------- How it works ----------
    about.column_dimensions["A"].width = 30
    about.column_dimensions["B"].width = 90
    about["A1"] = "How the prediction works"
    about["A1"].font = TITLE
    text = [
        ("What it is", "Rules, not a trained model: the chance each palm carries tender nuts (the harvest target) "
                       "right now, from the latest evidence in the ground station's database."),
        ("Bunch cycle", f"A palm opens a new bunch about every {ripeness.BUNCH_INTERVAL_D:.0f} days; a bunch stays "
                        f"tender for about {ripeness.TENDER_WINDOW_D:.0f} days."),
        ("Confirmed palm", "100%: a confirm stop already saw tender nuts."),
        ("Tender sighting", f"Its confidence (survey 'ripe' {ripeness.SEED_TENDER_CONF:.0%}, a logged visit "
                            f"{ripeness.VISIT_CONF:.0%}, ML or colour card: their own), fading after the tender window."),
        ("Mature / over-mature", f"{ripeness.STAGE_P['mature']:.0%} / {ripeness.STAGE_P['over-mature']:.0%}: "
                                 "a bunch was missed, a younger one is likely tender now."),
        ("Young", f"{ripeness.STAGE_P['young']:.0%}: tender in a few months, not now."),
        ("Harvested", "Due again about one bunch interval after the harvest."),
        ("Never observed", f"{ripeness.UNKNOWN_P:.0%}."),
        ("Health", f"x {ripeness.HEALTH_FACTOR} for frond damage or pests (stressed palms bear less)."),
        ("Predicted tender", f"Probability at or above {ripeness.PREDICT_MIN:.0%}; these join the next mission "
                             "after the targets, best first, while the battery stays within its limit."),
        ("Source", "backend/ripeness.py and backend/planner.py in the ground station."),
    ]
    for r, (k, v) in enumerate(text, start=3):
        about.cell(row=r, column=1, value=k).font = BOLD
        c = about.cell(row=r, column=2, value=v)
        c.font = BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
        about.cell(row=r, column=1).alignment = Alignment(vertical="top")

    wb.calculation.fullCalcOnLoad = True           # Excel computes the formulas when the file opens
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
