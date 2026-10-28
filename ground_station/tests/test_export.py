"""Prediction export to Excel (backend/export_xlsx.py, GET /api/prediction.xlsx)."""
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app


def test_prediction_xlsx():
    with TestClient(create_app()) as client:
        r = client.get("/api/prediction.xlsx")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
        assert "attachment; filename=prediction_" in r.headers["content-disposition"]
        trees = client.get("/api/trees").json()
        plan = client.get("/api/plan").json()
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Prediction", "Next mission", "How it works"]
    p = wb["Prediction"]
    assert p.max_row - 1 == len(trees)
    scores = [p.cell(i, 4).value for i in range(2, p.max_row + 1)]
    assert scores == sorted(scores, reverse=True)                       # most likely tender first
    assert p["E2"].value == '=IF(D2>=Summary!$C$5,"Yes","No")'          # live against the threshold
    by_id = {t["id"]: t for t in trees}
    assert abs(p["D2"].value - by_id[p["A2"].value]["ripe_score"]) < 1e-9
    in_mission = {p.cell(i, 1).value for i in range(2, p.max_row + 1) if p.cell(i, 10).value == "Yes"}
    assert in_mission == {s["tree"] for s in plan["stops"]}
    s = wb["Summary"]
    assert s["C5"].value == 0.5 and s["C8"].value.startswith("=COUNTA(Prediction!")
    m = wb["Next mission"]
    assert [m.cell(i, 2).value for i in range(2, 2 + len(plan["stops"]))] == [x["tree"] for x in plan["stops"]]
